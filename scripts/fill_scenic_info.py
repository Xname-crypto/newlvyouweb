from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import sys
import time
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Iterable
from urllib.parse import urlparse

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = BASE_DIR / "旅游景点.xlsx"
DEFAULT_OUTPUT = BASE_DIR / "旅游景点_补全.xlsx"
DEFAULT_CACHE_DIR = BASE_DIR / ".codex-cache" / "scenic-info"

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

HEADERS = {
    "city": ("城市",),
    "name": ("名称",),
    "booking": ("是否需要预定", "是否需要预约", "是否预约", "预约要求"),
    "opening": ("开放时间", "营业时间", "开放时段", "开馆时间"),
}

SEARCH_DOMAINS = {
    "bing.com",
    "cn.bing.com",
    "baidu.com",
    "www.baidu.com",
    "m.baidu.com",
    "sogou.com",
    "www.sogou.com",
    "so.com",
    "sm.cn",
}

LOW_VALUE_DOMAINS = {
    "zhihu.com",
    "www.zhihu.com",
    "xiaohongshu.com",
    "www.xiaohongshu.com",
    "bilibili.com",
    "www.bilibili.com",
    "douyin.com",
    "www.douyin.com",
    "weibo.com",
    "www.weibo.com",
    "tieba.baidu.com",
    "dahepiao.com",
    "www.dahepiao.com",
    "baike.baidu.com",
}

PLATFORM_DOMAINS = {
    "ctrip.com",
    "trip.com",
    "you.ctrip.com",
    "vacations.ctrip.com",
    "ly.com",
    "meituan.com",
    "qunar.com",
    "mafengwo.cn",
    "fliggy.com",
    "dianping.com",
}

SKIP_EXTENSIONS = {
    ".jpg",
    ".jpeg",
    ".png",
    ".gif",
    ".webp",
    ".svg",
    ".pdf",
    ".doc",
    ".docx",
    ".xls",
    ".xlsx",
    ".ppt",
    ".pptx",
    ".zip",
}

ANTI_BOT_MARKERS = (
    "访问过于频繁",
    "系统检测到",
    "请输入验证码",
    "Access Denied",
    "安全验证",
    "robot or human",
)

OPENING_LABELS = (
    "开放时间",
    "营业时间",
    "开放时段",
    "开馆时间",
    "开园时间",
    "运营时间",
    "对外开放时间",
    "游玩时间",
    "参观时间",
)

OPENING_NOISE = (
    "服务时间",
    "客服",
    "咨询电话",
    "退票",
    "热线",
    "停车",
    "地址",
    "票价",
    "门票",
    "购票",
)

BOOKING_POSITIVE = (
    "需要预约",
    "需预约",
    "须预约",
    "提前预约",
    "实名预约",
    "分时预约",
    "线上预约",
    "预约入园",
    "预约参观",
    "预约购票",
    "实名制分时段购票",
    "线上实名制分时段购票",
    "请提前预约",
    "所有观众均须",
    "均须通过官方渠道",
    "需检票后方可免费预约",
    "实名制购票政策",
    "实名制购票",
    "门票预订",
    "立即预订",
)

BOOKING_NEGATIVE = (
    "无需预约",
    "无须预约",
    "不需要预约",
    "免预约",
    "可直接入园",
    "直接入园",
    "现场购票入园",
)

TIME_RANGE_RE = re.compile(
    r"(?P<start_prefix>上午|下午|晚上|中午|凌晨|AM|PM|am|pm)?\s*"
    r"(?P<start>(?:[01]?\d|2[0-3])[:：][0-5]\d)\s*"
    r"(?:-|–|—|~|～|至|到)\s*"
    r"(?P<end_prefix>上午|下午|晚上|中午|凌晨|AM|PM|am|pm)?\s*"
    r"(?P<end>(?:[01]?\d|2[0-3])[:：][0-5]\d)"
)
DAY_RANGE_RE = re.compile(
    r"(周[一二三四五六日天末]|星期[一二三四五六日天末])(?:\s*(?:至|到|-|—|~|～)\s*(周[一二三四五六日天末]|星期[一二三四五六日天末]))?"
)


@dataclass
class SearchResult:
    url: str
    title: str
    snippet: str


@dataclass
class ScenicInfo:
    booking: str = "未知"
    opening: str = ""
    source_url: str = ""
    source_title: str = ""


class ScenicInfoFiller:
    def __init__(self, args: argparse.Namespace) -> None:
        self.args = args
        self.input_path = Path(args.input).resolve()
        self.output_path = Path(args.output).resolve()
        self.cache_dir = Path(args.cache_dir).resolve()
        self.cache_dir.mkdir(parents=True, exist_ok=True)

        self.session = requests.Session()
        self.session.headers.update(
            {
                "User-Agent": (
                    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                    "AppleWebKit/537.36 (KHTML, like Gecko) "
                    "Chrome/124.0.0.0 Safari/537.36"
                ),
                "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
                "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
            }
        )
        self._last_request_ts = 0.0

    def run(self) -> None:
        if not self.input_path.exists() and not self.output_path.exists():
            raise FileNotFoundError(f"Input file not found: {self.input_path}")

        workbook = self._prepare_workbook()
        sheet = workbook.active
        columns = self._ensure_columns(sheet)
        total_rows = sheet.max_row

        processed = 0
        updated = 0

        for row_idx in range(self.args.start_row, total_rows + 1):
            if self.args.limit is not None and processed >= self.args.limit:
                break

            city = self._cell_text(sheet.cell(row=row_idx, column=columns["city"]).value)
            name = self._cell_text(sheet.cell(row=row_idx, column=columns["name"]).value)
            booking_cell = sheet.cell(row=row_idx, column=columns["booking"])
            opening_cell = sheet.cell(row=row_idx, column=columns["opening"])
            current_booking = self._cell_text(booking_cell.value)
            current_opening = self._cell_text(opening_cell.value)

            processed += 1

            if not name:
                print(f"[skip] row={row_idx} empty scenic name")
                continue

            if not self.args.overwrite and current_booking and current_opening:
                print(f"[skip] row={row_idx} {name} already filled")
                continue

            print(f"[scan] row={row_idx} city={city or '-'} name={name}")
            try:
                result = self._resolve_scenic_info(city=city, name=name)
            except Exception as exc:
                print(f"[row-fail] row={row_idx} name={name} err={exc}")
                result = ScenicInfo()

            changed = False
            if self.args.overwrite or not current_booking:
                new_booking = result.booking or "未知"
                if new_booking != current_booking:
                    booking_cell.value = new_booking
                    changed = True

            if self.args.overwrite or not current_opening:
                new_opening = result.opening or ""
                if new_opening != current_opening:
                    opening_cell.value = new_opening
                    changed = True

            if changed:
                updated += 1
                print(
                    "[write] row={row} booking={booking} opening={opening} source={source}".format(
                        row=row_idx,
                        booking=booking_cell.value or "",
                        opening=opening_cell.value or "",
                        source=result.source_url or "",
                    )
                )

            if processed % self.args.checkpoint_every == 0:
                workbook.save(self.output_path)
                print(f"[save] checkpoint rows={processed} updated={updated} file={self.output_path}")

        workbook.save(self.output_path)
        print(f"[done] processed={processed} updated={updated} output={self.output_path}")

    def _prepare_workbook(self):
        if self.output_path.exists():
            print(f"[resume] loading existing output workbook: {self.output_path}")
        else:
            source = self.input_path
            print(f"[init] copying workbook from {source} to {self.output_path}")
            shutil.copy2(source, self.output_path)

        return load_workbook(self.output_path)

    def _ensure_columns(self, sheet) -> dict[str, int]:
        columns: dict[str, int] = {}
        for key, aliases in HEADERS.items():
            columns[key] = self._find_or_create_column(sheet, key, aliases)
        return columns

    def _find_or_create_column(self, sheet, key: str, aliases: Iterable[str]) -> int:
        alias_set = {alias.strip() for alias in aliases}
        for col_idx in range(1, sheet.max_column + 1):
            value = self._cell_text(sheet.cell(row=1, column=col_idx).value)
            if value in alias_set:
                return col_idx

        if key in {"city", "name"}:
            raise KeyError(f"Missing required source column: {aliases[0]}")

        col_idx = sheet.max_column + 1
        sheet.cell(row=1, column=col_idx).value = aliases[0]
        print(f"[column] created {aliases[0]} at index {col_idx}")
        return col_idx

    def _resolve_scenic_info(self, city: str, name: str) -> ScenicInfo:
        cache_path = self._cache_path(city, name)
        if cache_path.exists() and not self.args.refresh_cache:
            try:
                payload = json.loads(cache_path.read_text(encoding="utf-8"))
                return ScenicInfo(**payload)
            except Exception:
                pass

        queries = self._build_queries(city=city, name=name)
        fetched_urls: set[str] = set()
        best = ScenicInfo()

        for query in queries:
            try:
                results = self._search_candidates(query)
            except Exception as exc:
                print(f"[search-fail] query={query} err={exc}")
                continue
            for result in self._rank_results(results, scenic_name=name):
                if result.url in fetched_urls:
                    continue
                fetched_urls.add(result.url)

                page_text = self._fetch_page_text(result.url)
                if not page_text:
                    continue

                opening = best.opening or self._extract_opening_time(page_text)
                booking = best.booking
                if booking == "未知":
                    extracted_booking = self._extract_booking_status(page_text)
                    if extracted_booking:
                        booking = extracted_booking

                if opening or booking != "未知":
                    best = ScenicInfo(
                        booking=booking or "未知",
                        opening=opening or "",
                        source_url=result.url,
                        source_title=result.title,
                    )

                if best.opening and best.booking != "未知":
                    self._write_cache(cache_path, best)
                    return best

        self._write_cache(cache_path, best)
        return best

    def _build_queries(self, city: str, name: str) -> list[str]:
        cleaned_city = self._clean_query_text(city)
        cleaned_name = self._clean_query_text(name)
        exact_name = f"\"{cleaned_name}\"" if cleaned_name else ""
        queries = [
            " ".join(part for part in (cleaned_city, exact_name, "开放时间 官方") if part),
            " ".join(part for part in (cleaned_city, exact_name, "预约 官方") if part),
            " ".join(part for part in ("site:ctrip.com", exact_name, "开放时间") if part),
            " ".join(part for part in ("site:ly.com", exact_name, "开放时间") if part),
            " ".join(part for part in ("site:meituan.com", exact_name, "开放时间") if part),
            " ".join(part for part in ("site:ctrip.com", exact_name, "预约") if part),
        ]
        return list(dict.fromkeys(query for query in queries if query))

    def _search_candidates(self, query: str) -> list[SearchResult]:
        return self._search_baidu(query)

    def _search_baidu(self, query: str) -> list[SearchResult]:
        self._sleep_between_requests()
        response = self.session.get(
            "https://www.baidu.com/s",
            params={"wd": query},
            timeout=self.args.timeout,
        )
        response.raise_for_status()

        results = self._parse_baidu_results(response.text)
        print(f"[search] engine=baidu query={query} results={len(results)}")
        return results[: self.args.max_results]

    def _search_sogou(self, query: str) -> list[SearchResult]:
        self._sleep_between_requests()
        response = self.session.get(
            "https://www.sogou.com/web",
            params={"query": query},
            timeout=self.args.timeout,
        )
        response.raise_for_status()

        results = self._parse_sogou_results(response.text)
        print(f"[search] engine=sogou query={query} results={len(results)}")
        return results[: self.args.max_results]

    def _parse_baidu_results(self, html: str) -> list[SearchResult]:
        soup = BeautifulSoup(html, "html.parser")
        parsed: list[SearchResult] = []
        seen: set[str] = set()

        for block in soup.select("div.result, div.result-op, div.c-container"):
            link = block.select_one("h3 a") or block.select_one("a")
            if not link:
                continue
            raw_url = link.get("href", "")
            url = self._resolve_search_result_url(raw_url)
            if not url or url in seen:
                continue

            title = self._clean_line(link.get_text(" ", strip=True))
            snippet_node = block.select_one("div.c-abstract") or block.select_one("div.content-right_8Zs40") or block.select_one("p")
            snippet = self._clean_line(snippet_node.get_text(" ", strip=True) if snippet_node else "")

            if self._should_skip_url(url):
                continue

            seen.add(url)
            parsed.append(SearchResult(url=url, title=title, snippet=snippet))

        return parsed

    def _parse_sogou_results(self, html: str) -> list[SearchResult]:
        soup = BeautifulSoup(html, "html.parser")
        parsed: list[SearchResult] = []
        seen: set[str] = set()

        for block in soup.select("div.vrwrap, div.rb, div.results > div"):
            link = block.select_one("h3 a") or block.select_one("a")
            if not link:
                continue
            raw_url = link.get("href", "")
            url = self._resolve_search_result_url(raw_url)
            if not url or url in seen:
                continue

            title = self._clean_line(link.get_text(" ", strip=True))
            snippet_node = block.select_one("div.text-layout, p, div.ft")
            snippet = self._clean_line(snippet_node.get_text(" ", strip=True) if snippet_node else "")

            if self._should_skip_url(url):
                continue

            seen.add(url)
            parsed.append(SearchResult(url=url, title=title, snippet=snippet))

        return parsed

    def _rank_results(self, results: list[SearchResult], scenic_name: str) -> list[SearchResult]:
        def score(item: SearchResult) -> tuple[int, int]:
            host = urlparse(item.url).netloc.lower()
            title_text = f"{item.title} {item.snippet}"
            value = 0

            if any(host == domain or host.endswith(f".{domain}") for domain in PLATFORM_DOMAINS):
                value += 25
            else:
                value += 45

            if host.endswith(".gov.cn") or host.endswith(".org.cn"):
                value += 12

            if "官方" in title_text or "公告" in title_text or "票务" in title_text:
                value += 10

            if scenic_name and scenic_name in title_text:
                value += 6

            if "开放时间" in title_text or "营业时间" in title_text:
                value += 5

            if "预约" in title_text or "购票" in title_text:
                value += 5

            if any(noise in title_text for noise in ("攻略", "游记", "问答", "论坛")):
                value -= 8

            return value, -len(item.url)

        return sorted(results, key=score, reverse=True)[: self.args.max_pages]

    def _fetch_page_text(self, url: str) -> str:
        for fetcher in (self._fetch_direct_text, self._fetch_jina_text):
            try:
                text = fetcher(url)
            except Exception as exc:
                print(f"[fetch-fail] url={url} method={fetcher.__name__} err={exc}")
                continue

            if text and not self._looks_blocked(text):
                print(f"[fetch] url={url} method={fetcher.__name__} chars={len(text)}")
                return text

        return ""

    def _fetch_direct_text(self, url: str) -> str:
        self._sleep_between_requests()
        response = self.session.get(url, timeout=self.args.timeout)
        response.raise_for_status()

        content_type = response.headers.get("Content-Type", "").lower()
        if "text/html" not in content_type and "application/xhtml+xml" not in content_type:
            return ""

        if not response.encoding or response.encoding.lower() == "iso-8859-1":
            response.encoding = response.apparent_encoding or "utf-8"

        soup = BeautifulSoup(response.text, "html.parser")
        for node in soup(["script", "style", "noscript", "iframe"]):
            node.extract()

        text = soup.get_text("\n")
        return self._normalize_text_block(text)

    def _fetch_jina_text(self, url: str) -> str:
        self._sleep_between_requests()
        response = self.session.get(f"https://r.jina.ai/{url}", timeout=self.args.timeout + 10)
        response.raise_for_status()
        return self._normalize_text_block(response.text)

    def _extract_booking_status(self, text: str) -> str | None:
        lines = self._split_lines(text)
        positive_hits = 0
        negative_hits = 0

        for line in lines:
            if not any(token in line for token in ("预约", "预定", "预订", "购票", "入园")):
                continue

            if any(marker in line for marker in BOOKING_NEGATIVE):
                negative_hits += 2

            if any(marker in line for marker in BOOKING_POSITIVE):
                positive_hits += 2

            if "实名制分时段购票" in line or "均须通过官方渠道" in line:
                positive_hits += 2

            if "实名制购票" in line or "门票预订" in line:
                positive_hits += 2

            if "立即预订" in line and ("门票" in line or "乐园" in line or "景区" in line):
                positive_hits += 2

            if "可直接入园" in line:
                negative_hits += 2

        if negative_hits > positive_hits and negative_hits >= 2:
            return "否"
        if positive_hits >= 2:
            return "是"
        return None

    def _extract_opening_time(self, text: str) -> str:
        lines = self._split_lines(text)
        candidates: list[tuple[int, str]] = []

        for idx, line in enumerate(lines):
            merged = self._merge_line_window(lines, idx)
            if not merged:
                continue

            score = 0
            has_label = any(label in merged for label in OPENING_LABELS)
            has_time = bool(TIME_RANGE_RE.search(merged))
            has_day = bool(DAY_RANGE_RE.search(merged))

            if has_label:
                score += 6
            if has_time:
                score += 5
            if has_day:
                score += 3
            if "闭馆" in merged or "停止入场" in merged:
                score += 2
            if any(noise in merged for noise in OPENING_NOISE) and not has_label:
                score -= 6

            if score <= 0:
                continue

            cleaned = self._cleanup_opening_candidate(merged)
            if cleaned:
                candidates.append((score, cleaned))

        if not candidates:
            return ""

        candidates.sort(key=lambda item: (item[0], -len(item[1])), reverse=True)
        return candidates[0][1]

    def _merge_line_window(self, lines: list[str], idx: int) -> str:
        line = lines[idx]
        if len(line) > 120:
            line = line[:120]

        chunks = [line]
        for offset in (1, 2):
            next_idx = idx + offset
            if next_idx >= len(lines):
                break
            next_line = lines[next_idx]
            if len(next_line) > 60:
                break
            if any(marker in next_line for marker in ("地址", "电话", "交通", "门票")):
                break
            chunks.append(next_line)

        return self._clean_line(" ".join(chunks))

    def _cleanup_opening_candidate(self, text: str) -> str:
        value = text
        value = re.sub(r"\s+", " ", value).strip(" ，,;；。")

        time_match = TIME_RANGE_RE.search(value)
        if not time_match:
            return ""

        label_positions = [value.rfind(label, 0, time_match.start()) for label in OPENING_LABELS]
        label_start = max((pos for pos in label_positions if pos >= 0), default=-1)
        context_start = label_start if label_start >= 0 else max(0, time_match.start() - 20)
        context = value[context_start:].strip()

        context = re.sub(r"^[^开放营业开馆开园运营参观游玩]{0,20}", "", context)
        label_regex = "|".join(re.escape(label) for label in OPENING_LABELS)
        context = re.sub(rf"^(?:{label_regex})\s*[:：]?\s*", "", context)
        context = context.strip()

        day_match = DAY_RANGE_RE.search(context)
        day_text = self._clean_line(day_match.group(0)) if day_match else ""
        start_text = self._normalize_clock(time_match.group("start_prefix"), time_match.group("start"))
        end_text = self._normalize_clock(time_match.group("end_prefix"), time_match.group("end"))
        base = f"{start_text}-{end_text}"
        if day_text:
            base = f"{day_text} {base}"

        closing_match = re.search(r"(周[一二三四五六日天末][^，。；]{0,8}闭馆|闭馆日[^，。；]{0,12}|周一闭馆)", value)
        if closing_match and closing_match.group(0) not in base:
            base = f"{base}，{self._clean_line(closing_match.group(0))}"

        return base.strip(" ，,;；。")

    def _normalize_clock(self, prefix: str | None, value: str) -> str:
        clean = value.replace("：", ":").strip()
        hour_text, minute_text = clean.split(":")
        hour = int(hour_text)
        minute = int(minute_text)

        if prefix in {"下午", "晚上", "PM", "pm"} and hour < 12:
            hour += 12
        elif prefix == "中午" and hour < 11:
            hour += 12
        elif prefix == "凌晨" and hour == 12:
            hour = 0

        return f"{hour:02d}:{minute:02d}"

    def _sleep_between_requests(self) -> None:
        if self.args.request_interval <= 0:
            return

        elapsed = time.time() - self._last_request_ts
        if elapsed < self.args.request_interval:
            time.sleep(self.args.request_interval - elapsed)
        self._last_request_ts = time.time()

    def _cache_path(self, city: str, name: str) -> Path:
        key = hashlib.md5(f"{city}|{name}".encode("utf-8")).hexdigest()
        return self.cache_dir / f"{key}.json"

    def _write_cache(self, path: Path, info: ScenicInfo) -> None:
        path.write_text(
            json.dumps(asdict(info), ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

    def _normalize_candidate_url(self, value: str) -> str:
        value = value.strip()
        if not value or value.startswith("javascript:"):
            return ""
        if value.startswith("//"):
            value = f"https:{value}"
        if value.startswith("/"):
            value = f"https://www.sogou.com{value}"
        if not value.startswith("http"):
            return ""
        return value

    def _resolve_search_result_url(self, value: str) -> str:
        url = self._normalize_candidate_url(value)
        if not url:
            return ""

        if "baidu.com/baidu.php?" in url:
            return ""

        if any(marker in url for marker in ("baidu.com/link?", "sogou.com/link?")):
            try:
                self._sleep_between_requests()
                response = self.session.get(url, timeout=self.args.timeout, allow_redirects=True, stream=True)
                resolved = response.url
                response.close()
                return self._normalize_candidate_url(resolved)
            except Exception as exc:
                print(f"[resolve-fail] url={url} err={exc}")
                return ""

        return url

    def _should_skip_url(self, url: str) -> bool:
        parsed = urlparse(url)
        host = parsed.netloc.lower()
        path = parsed.path.lower()

        if any(host == domain or host.endswith(f".{domain}") for domain in SEARCH_DOMAINS):
            return True
        if any(host == domain or host.endswith(f".{domain}") for domain in LOW_VALUE_DOMAINS):
            return True
        if any(path.endswith(ext) for ext in SKIP_EXTENSIONS):
            return True
        return False

    def _looks_blocked(self, text: str) -> bool:
        if len(text) < 80:
            return True
        return any(marker.lower() in text.lower() for marker in ANTI_BOT_MARKERS)

    def _split_lines(self, text: str) -> list[str]:
        lines: list[str] = []
        for raw in text.splitlines():
            line = self._clean_line(raw)
            if line:
                lines.append(line)
        return lines

    def _normalize_text_block(self, text: str) -> str:
        translation = str.maketrans(
            {
                "\u3000": " ",
                "\xa0": " ",
                "：": ":",
                "（": "(",
                "）": ")",
            }
        )
        normalized = text.translate(translation)
        return normalized

    def _clean_line(self, value: str) -> str:
        value = self._normalize_text_block(value)
        value = re.sub(r"\s+", " ", value)
        return value.strip()

    def _cell_text(self, value) -> str:
        if value is None:
            return ""
        return str(value).strip()

    def _clean_query_text(self, value: str) -> str:
        value = self._cell_text(value)
        value = re.sub(r"[()（）【】\[\]]", " ", value)
        value = re.sub(r"\s+", " ", value)
        return value.strip()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Fill scenic booking and opening info into Excel.")
    parser.add_argument("--input", default=str(DEFAULT_INPUT), help="Path to input workbook.")
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT), help="Path to output workbook.")
    parser.add_argument("--cache-dir", default=str(DEFAULT_CACHE_DIR), help="Directory for scenic cache files.")
    parser.add_argument("--limit", type=int, default=None, help="Process only the first N rows from start-row.")
    parser.add_argument("--start-row", type=int, default=2, help="Starting row number in Excel.")
    parser.add_argument("--checkpoint-every", type=int, default=10, help="Save workbook every N processed rows.")
    parser.add_argument("--max-results", type=int, default=8, help="Max search results to parse for each query.")
    parser.add_argument("--max-pages", type=int, default=3, help="Max candidate pages to fetch per query.")
    parser.add_argument("--request-interval", type=float, default=1.2, help="Delay between outbound requests in seconds.")
    parser.add_argument("--timeout", type=int, default=20, help="HTTP timeout in seconds.")
    parser.add_argument("--overwrite", action="store_true", help="Overwrite existing filled values in output workbook.")
    parser.add_argument("--refresh-cache", action="store_true", help="Ignore existing scenic cache and refetch.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    ScenicInfoFiller(args).run()


if __name__ == "__main__":
    main()
